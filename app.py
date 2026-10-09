"""Asistente Pedagógico Perú V1.0. API Flask para sesiones docentes."""
import io
import json
import os
import re
from datetime import datetime
from flask import Flask, jsonify, render_template, request, send_file
from dotenv import load_dotenv
from docx import Document
from docx.shared import Cm, Pt, RGBColor

load_dotenv()
app = Flask(__name__)
app.json.ensure_ascii = False

# Evitar que el navegador reutilice JavaScript de versiones anteriores.
@app.after_request
def evitar_cache_frontend(respuesta):
    if request.path == '/' or request.path.endswith(('.js', '.css', '.html')):
        respuesta.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        respuesta.headers['Pragma'] = 'no-cache'
        respuesta.headers['Expires'] = '0'
    return respuesta


COMPETENCIAS = {
    'Comunicación': ['Se comunica oralmente en su lengua materna', 'Lee diversos tipos de textos escritos en su lengua materna', 'Escribe diversos tipos de textos en su lengua materna'],
    'Matemática': ['Resuelve problemas de cantidad', 'Resuelve problemas de regularidad, equivalencia y cambio', 'Resuelve problemas de forma, movimiento y localización', 'Resuelve problemas de gestión de datos e incertidumbre'],
    'Personal Social': ['Construye su identidad', 'Convive y participa democráticamente en la búsqueda del bien común', 'Construye interpretaciones históricas', 'Gestiona responsablemente el espacio y el ambiente', 'Gestiona responsablemente los recursos económicos'],
    'Ciencia y Tecnología': ['Indaga mediante métodos científicos para construir sus conocimientos', 'Explica el mundo físico basándose en conocimientos sobre los seres vivos, materia y energía, biodiversidad, Tierra y universo', 'Diseña y construye soluciones tecnológicas para resolver problemas de su entorno'],
    'Arte y Cultura': ['Aprecia de manera crítica manifestaciones artístico-culturales', 'Crea proyectos desde los lenguajes artísticos'],
    'Educación Física': ['Se desenvuelve de manera autónoma a través de su motricidad', 'Asume una vida saludable', 'Interactúa a través de sus habilidades sociomotrices'],
    'Educación Religiosa': ['Construye su identidad como persona humana, amada por Dios, digna, libre y trascendente, comprendiendo la doctrina de su propia religión, abierto al diálogo con las que le son cercanas', 'Asume la experiencia del encuentro personal y comunitario con Dios en su proyecto de vida en coherencia con su creencia religiosa'],
    'Inglés': ['Se comunica oralmente en inglés como lengua extranjera', 'Lee diversos tipos de textos escritos en inglés como lengua extranjera', 'Escribe diversos tipos de textos en inglés como lengua extranjera'],
    'Tutoría': []
}
ENFOQUES = ['De derechos', 'Inclusivo o de atención a la diversidad', 'Intercultural', 'Igualdad de género', 'Ambiental', 'Orientación al bien común', 'Búsqueda de la excelencia']
FUENTES = [
    'https://www.gob.pe/institucion/minedu/informes-publicaciones/6790803-programa-curricular-de-educacion-primaria',
    'https://minedu.gob.pe/curriculo/recursos.php',
    'https://repositorio.minedu.gob.pe/handle/20.500.12799/5310'
]


def sanitizar(d):
    if not isinstance(d, dict):
        raise ValueError('Se esperaba un formulario válido.')
    campos = ['docente', 'institucion', 'grado', 'area', 'tema', 'duracion', 'competencia', 'contexto', 'enfoque', 'necesidades', 'indicaciones']
    return {k: str(d.get(k, '') or '').strip()[:1500] for k in campos}


def propuesta_local(d):
    tema = d['tema']; grado = d['grado']; area = d['area']; duracion = int(d['duracion'])
    inicio = max(10, round(duracion * .17)); cierre = max(10, round(duracion * .17)); desarrollo = duracion - inicio - cierre
    competencia = d['competencia'] or 'Por seleccionar y verificar en el Programa Curricular de Primaria'
    foco = d['contexto'] or 'situaciones de la vida cotidiana de la comunidad'
    return {
        'titulo': f'Exploramos {tema}', 'datos': d,
        'proposito': f'Que las y los estudiantes de {grado}.° grado movilicen sus saberes para comprender y aplicar {tema} en {foco}, mediante una actividad concreta y una explicación de lo aprendido.',
        'competencia': competencia,
        'capacidades': 'Pendiente: seleccionar las capacidades oficiales asociadas a la competencia en el Programa Curricular de Primaria.',
        'desempeno': 'Pendiente de validación: seleccionar el desempeño del grado que corresponda al propósito y al contexto.',
        'criterios': [f'Explica con sus palabras las ideas principales relacionadas con {tema}.', f'Aplica lo aprendido sobre {tema} para resolver una situación contextualizada.', 'Justifica su respuesta o producción con ejemplos comprensibles.'],
        'evidencia': f'Producción individual o grupal sobre {tema}, acompañada de una explicación oral o escrita.',
        'instrumento': 'Lista de cotejo basada en los criterios de evaluación.',
        'enfoque': d['enfoque'] or 'Orientación al bien común',
        'inicio': {'minutos': inicio, 'actividades': f'Acogida y presentación de una situación vinculada a {foco}. Preguntar: ¿qué sabemos sobre {tema}?, ¿dónde lo observamos?, ¿qué necesitamos averiguar? Recoger saberes previos sin calificarlos. Compartir propósito y criterios en lenguaje accesible.'},
        'desarrollo': {'minutos': desarrollo, 'actividades': f'Plantear un reto concreto acerca de {tema} adecuado a {grado}.° grado. En parejas o equipos, explorar materiales o información pertinente, registrar hallazgos y contrastar estrategias. El docente formula preguntas de profundización, acompaña a quienes necesitan apoyo y ofrece retroalimentación según los criterios. Cada equipo presenta su producción y explica cómo resolvió el reto.'},
        'cierre': {'minutos': cierre, 'actividades': 'Socializar conclusiones, revisar las evidencias frente a los criterios, realizar autoevaluación y responder: ¿qué aprendí?, ¿cómo lo aprendí?, ¿para qué lo usaré? Proponer una aplicación breve para casa o comunidad.'},
        'recursos': ['Pizarra y plumones', 'Hojas o cuaderno de trabajo', 'Material concreto, textos o imágenes relacionados con el tema'],
        'atencion_diversidad': d['necesidades'] or 'Ofrecer distintas formas de participación, apoyos visuales, consignas breves y tiempo flexible, según necesidades observadas.',
        'retroalimentacion': 'Describir el avance respecto de los criterios, preguntar por las decisiones tomadas y acordar una mejora específica sin limitarse a asignar una nota.',
        'lista_cotejo': [{'criterio': f'Identifica ideas centrales sobre {tema}', 'si': '', 'no': ''}, {'criterio': 'Aplica lo aprendido en el reto', 'si': '', 'no': ''}, {'criterio': 'Explica o justifica su procedimiento', 'si': '', 'no': ''}],
        'observaciones': ['Borrador orientativo: NO constituye una sesión aprobada por MINEDU.', 'Las capacidades y desempeños oficiales deben comprobarse en la fuente curricular antes de usarlo.'],
        'origen': 'plantilla_local'
    }


def generar_ia(d):
    clave = os.getenv('GROQ_API_KEY', '').strip()
    if not clave:
        return propuesta_local(d)
    from openai import OpenAI
    cliente = OpenAI(api_key=clave, base_url='https://api.groq.com/openai/v1', timeout=45)
    modelo = os.getenv('GROQ_MODEL', 'llama-3.3-70b-versatile')
    base = propuesta_local(d)
    sistema = ("Eres asesor de planificación pedagógica de educación primaria peruana. "
               "Elabora sesiones prácticas según el enfoque por competencias del CNEB. "
               "Responde exclusivamente un objeto JSON con LAS MISMAS claves y tipos que el esquema dado. "
               "No inventes citas, números de norma, desempeños ni capacidades oficiales. Si no puedes corroborar un desempeño exacto, escribe 'Pendiente de validación' y una propuesta orientativa claramente identificada. "
               "Criterios observables, evidencia alineada, inclusión y evaluación formativa. "
               "Las actividades deben ser concretas y apropiadas al grado y tema; respetar los minutos. "
               "Evita confundir procesos pedagógicos generales con procesos didácticos específicos de área. "
               "No incluyas datos sensibles de menores.")
    resp = cliente.chat.completions.create(model=modelo, temperature=.45, max_tokens=3200,
      messages=[{'role':'system','content':sistema}, {'role':'user','content':'Datos de planificación: '+json.dumps(d, ensure_ascii=False)+'\nEsquema requerido: '+json.dumps(base, ensure_ascii=False)}],
      response_format={'type':'json_object'})
    obj = json.loads(resp.choices[0].message.content)
    for key, val in base.items():
        if key not in obj or not isinstance(obj[key], type(val)):
            obj[key] = val
    obj['datos'] = d
    obj['origen'] = 'ia_groq_borrador'
    for k in ('inicio','desarrollo','cierre'):
        if not isinstance(obj[k],dict): obj[k] = base[k]
    # Los minutos los calcula la aplicación, no el modelo.
    for k in ('inicio','desarrollo','cierre'):
        obj[k]['minutos'] = base[k]['minutos']
    obj['observaciones'] = list(obj.get('observaciones', []))[:8] + ['Verificar la exactitud curricular antes del uso docente.']
    return obj


@app.get('/')
def inicio():
    return render_template('index.html')


@app.get('/api/catalogo')
def catalogo():
    return jsonify({'competencias': COMPETENCIAS, 'enfoques': ENFOQUES, 'fuentes': FUENTES, 'modo_ia': bool(os.getenv('GROQ_API_KEY'))})


@app.post('/api/generar')
def generar():
    try:
        d = sanitizar(request.get_json(silent=True))
        if not d['tema'] or d['area'] not in COMPETENCIAS or d['grado'] not in [str(i) for i in range(1, 7)]:
            return jsonify({'error': 'Indique un tema, área y grado válidos.'}), 400
        if d['competencia'] and d['area'] != 'Tutoría' and d['competencia'] not in COMPETENCIAS[d['area']]:
            return jsonify({'error': 'La competencia no pertenece al área seleccionada.'}), 400
        try: dur = int(d['duracion'])
        except ValueError: dur = 90
        if not 30 <= dur <= 240:
            return jsonify({'error': 'La duración debe estar entre 30 y 240 minutos.'}), 400
        d['duracion'] = str(dur)
        try:
            sesion = generar_ia(d)
            return jsonify({'sesion': sesion, 'advertencia': None})
        except Exception as exc:
            app.logger.warning('IA no disponible: %s', type(exc).__name__)
            return jsonify({'sesion': propuesta_local(d), 'advertencia': 'La IA no respondió; se utilizó la plantilla local orientativa.'})
    except Exception:
        return jsonify({'error': 'No se pudo procesar el formulario.'}), 400


@app.post('/api/exportar-word')
def exportar_word():
    data = request.get_json(silent=True)
    if not isinstance(data, dict) or len(json.dumps(data)) > 100000:
        return jsonify({'error': 'Documento inválido o demasiado grande.'}), 400
    doc = Document()
    section = doc.sections[0]; section.top_margin = Cm(2); section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.4); section.right_margin = Cm(2.4)
    styles = doc.styles; styles['Normal'].font.name = 'Calibri'; styles['Normal'].font.size = Pt(10)
    styles['Normal'].paragraph_format.space_after = Pt(6)
    for s in ('Title','Heading 1','Heading 2'):
        styles[s].font.name = 'Calibri'; styles[s].font.color.rgb = RGBColor(25, 65, 103)
    doc.add_heading('SESIÓN DE APRENDIZAJE', 0)
    doc.add_paragraph(str(data.get('titulo','Sin título')))
    meta = data.get('datos', {}) if isinstance(data.get('datos'),dict) else {}
    table = doc.add_table(rows=0, cols=2); table.style='Light Shading Accent 1'
    for k, nombre in [('institucion','Institución educativa'),('docente','Docente'),('grado','Grado'),('area','Área'),('tema','Tema'),('duracion','Duración (min)')]:
        c = table.add_row().cells; c[0].text = nombre; c[1].text = str(meta.get(k,''))
    for titulo, key in [('Propósito de aprendizaje','proposito'),('Competencia','competencia'),('Capacidades (verificar)','capacidades'),('Desempeño (verificar)','desempeno'),('Evidencia','evidencia'),('Instrumento','instrumento'),('Enfoque transversal','enfoque')]:
        doc.add_heading(titulo, 2); doc.add_paragraph(str(data.get(key,'')))
    doc.add_heading('Criterios de evaluación', 2)
    for item in data.get('criterios',[])[:15]: doc.add_paragraph(str(item),style='List Bullet')
    doc.add_heading('Secuencia didáctica', 1)
    for clave, titulo in [('inicio','Inicio'),('desarrollo','Desarrollo'),('cierre','Cierre')]:
        info = data.get(clave,{})
        if not isinstance(info,dict): info={}
        doc.add_heading(f'{titulo} — {info.get("minutos", "?")} minutos', 2)
        doc.add_paragraph(str(info.get('actividades','')))
    for titulo, key in [('Recursos','recursos'),('Atención a la diversidad','atencion_diversidad'),('Retroalimentación formativa','retroalimentacion')]:
        doc.add_heading(titulo, 2)
        val = data.get(key,[] if key=='recursos' else '')
        if isinstance(val,list):
            for x in val[:20]: doc.add_paragraph(str(x),style='List Bullet')
        else: doc.add_paragraph(str(val))
    doc.add_heading('Lista de cotejo orientativa', 2)
    tabla=doc.add_table(rows=1, cols=3); tabla.style='Table Grid'
    for i,n in enumerate(['Criterio','Sí','No']): tabla.rows[0].cells[i].text=n
    for x in data.get('lista_cotejo',[])[:15]:
        c=tabla.add_row().cells; c[0].text=str(x.get('criterio','')) if isinstance(x,dict) else str(x); c[1].text=''; c[2].text=''
    doc.add_heading('Revisión docente', 2)
    doc.add_paragraph('Documento generado como borrador de apoyo. Verifique desempeños, capacidades y coherencia curricular con fuentes oficiales del MINEDU.')
    for fuente in FUENTES: doc.add_paragraph(fuente)
    stream=io.BytesIO(); doc.save(stream); stream.seek(0)
    return send_file(stream, as_attachment=True, download_name='sesion_primaria_peru.docx', mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


if __name__ == '__main__':
    app.run(host='127.0.0.1', port=5000, debug=False)
